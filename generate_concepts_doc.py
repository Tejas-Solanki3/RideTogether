from docx import Document

document = Document()

document.add_heading('RideTogether - Concepts & Deployment Guide', 0)

document.add_heading('Core Concepts Used', level=1)

document.add_heading('Flutter & Dart', level=2)
document.add_paragraph('Location: Across the entire codebase.')
document.add_paragraph('The foundational framework and language used to build the cross-platform mobile application, providing a single codebase for both iOS and Android.')

document.add_heading('Riverpod (State Management)', level=2)
document.add_paragraph('Location: lib/state/providers.dart')
document.add_paragraph("Used to efficiently manage the app's global state, such as streaming real-time rides and managing the active user session without tightly coupling logic to the UI.")

document.add_heading('Firebase Authentication', level=2)
document.add_paragraph('Location: lib/data/firestore_repository.dart')
document.add_paragraph('Handles secure user registration and login, including the automatic dispatch of email verification links to ensure only valid students can access the app.')

document.add_heading('Cloud Firestore (NoSQL Database)', level=2)
document.add_paragraph('Location: lib/data/firestore_repository.dart & firestore.rules')
document.add_paragraph('A real-time cloud database used to store users, ride offers, and matches. It uses Security Rules to ensure users can only post rides under their own accounts.')

document.add_heading('Form Validation', level=2)
document.add_paragraph('Location: lib/ui/screens/login.dart & lib/ui/screens/post_ride.dart')
document.add_paragraph('Ensures data integrity before submitting to the backend, such as strictly enforcing that email addresses end with a verified campus domain (.edu or .ac.in).')

document.add_heading('Deployment Guide', level=1)

document.add_heading('Google Play Store (Android)', level=2)
document.add_paragraph('Requirements: A Google Play Developer account ($25 one-time) and a generated keystore file to digitally sign your application.')
document.add_paragraph('Command: Run `flutter build appbundle` in the terminal to generate the final .aab file, which you then upload to the Google Play Console.')

document.add_heading('Apple App Store (iOS)', level=2)
document.add_paragraph('Requirements: An Apple Developer account ($99/year), Xcode installed on a Mac, and the appropriate App ID and provisioning profiles configured.')
document.add_paragraph('Command: Run `flutter build ipa` in the terminal to generate the iOS build archive, and then use Xcode or Transporter to upload it to App Store Connect.')

document.save('RideTogether_Concepts_And_Deployment.docx')
print('Document generated.')
