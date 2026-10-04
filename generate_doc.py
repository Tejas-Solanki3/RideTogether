from docx import Document
from docx.shared import Inches, Pt

document = Document()

document.add_heading('RideTogether (Campus Carpool Matching App)', 0)

document.add_heading('Problem Statement', level=1)
document.add_paragraph('RideTogether wants an app where a logged-in student posts a ride offer or a ride request between two campus-area points, and the app matches and lists compatible rides so riders and drivers can coordinate. (With Proper Justification)')

document.add_heading('Objectives Achieved', level=1)
document.add_paragraph('• UI/Widgets: Designed Post Ride, Find Ride, and My Matches screens using Form, ListView and Card widgets.')
document.add_paragraph('• Styling/Theming: Applied Material 3 theming with clear driver/rider role indicators (Monochrome aesthetic).')
document.add_paragraph('• Dart Logic: Modeled Ride as a Dart class and used Riverpod to manage and expose the live list of matching rides.')
document.add_paragraph('• Figma: Designed a clean, guided flow covering every screen listed above with progress cues for the user.')

document.add_heading('Technical Documentation', level=1)

document.add_heading('Dart Logic & State Management (Riverpod)', level=2)
document.add_paragraph('The application leverages Riverpod for robust and reactive state management. The core models (e.g., `Ride` and `Student`) are defined in `lib/domain/models.dart`. The repositories (`DemoRideRepository` and `FirestoreRideRepository`) implement the data layer, and `lib/state/providers.dart` defines providers such as `ridesProvider` and `matchesProvider` to seamlessly sync the backend data with the UI.')
document.add_paragraph('The seat count on a ride offer dynamically decreases as riders join, which is instantly reflected live across the app via Riverpod streams attached to Firestore snapshot listeners.')

document.add_heading('Firestore Integration', level=2)
document.add_paragraph('The app uses Cloud Firestore to store user profiles (under `students` collection), ride offers/requests (under `rides` collection), and ride matches (`matches`). The security rules are configured to ensure students can only publish rides from their own verified accounts, guaranteeing a safe campus environment.')

document.add_page_break()

document.add_heading('Screens & Flow', level=1)

document.add_heading('1. Registration & Verification Screen', level=2)
document.add_paragraph('File Location: lib/ui/screens/login.dart')
document.add_paragraph('[Insert Image Here: Login & Registration Screen]')
document.add_paragraph('A secure gateway for students to create an account, mandating a verified university email (e.g., .edu or .ac.in).')

document.add_heading('2. Find a Ride Screen', level=2)
document.add_paragraph('File Location: lib/ui/screens/find_ride.dart')
document.add_paragraph('[Insert Image Here: Find Ride Screen]')
document.add_paragraph('Displays a live, Riverpod-managed feed of available carpool options with precise route, timing, and available seat counts.')

document.add_heading('3. Post a Ride Screen', level=2)
document.add_paragraph('File Location: lib/ui/screens/post_ride.dart')
document.add_paragraph('[Insert Image Here: Post Ride Screen]')
document.add_paragraph('A guided form allowing verified students to post ride offers or requests between campus checkpoints.')

document.add_heading('4. My Matches Screen', level=2)
document.add_paragraph('File Location: lib/ui/screens/my_matches.dart')
document.add_paragraph('[Insert Image Here: My Matches Screen]')
document.add_paragraph('A dashboard listing rides the student has connected with, as either a driver or a rider, complete with status indicators.')

document.save('RideTogether_Project_Documentation.docx')
print('Document generated.')
